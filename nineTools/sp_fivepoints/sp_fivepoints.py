# -*- coding: utf-8 -*-
"""把本目录中的五点描述 Markdown 转为无表头的 xlsx。

直接运行 main() 即可处理本目录下所有 .md 文件。根据 Markdown 表头识别：
- 德语：每个产品一行，A~E 为德语第 1~5 点；输出 de_YYYYMMDD.xlsx。
- 法语：每个产品一行，A~E 为法语第 1~5 点；输出 fr_YYYYMMDD.xlsx。
- 德法双语：德语整块、空两行、法语整块；输出 de&fr_YYYYMMDD.xlsx。

同一天若有多份相同语言模式的 Markdown，会先报错，不覆盖其中一份的结果。
"""

import argparse
import io
import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook


BASE = os.path.dirname(os.path.abspath(__file__))
PRODUCT_RE = re.compile(r"^##\s*产品\s*(\S+)\s*$")
ROW_RE = re.compile(r"^\|\s*(\d+)\s*\|")


def language_of(label):
    """从表头单元格中识别语言，不依赖文件名或目录名。"""
    value = label.strip().casefold()
    if "德语" in value or "deutsch" in value or "german" in value:
        return "de"
    if "法语" in value or "français" in value or "francais" in value or "french" in value:
        return "fr"
    if re.match(r"^de(?:\b|[_:/-])", value):
        return "de"
    if re.match(r"^fr(?:\b|[_:/-])", value):
        return "fr"
    raise ValueError("无法从表头 %r 判断语言" % label)


def header_languages(cells, line_no):
    """返回表头中正文列的语言顺序，例如 ('de', 'fr')。"""
    if len(cells) not in (3, 5):
        raise ValueError("第 %d 行表头应为 3 列或 5 列" % line_no)
    languages = tuple(language_of(cells[i]) for i in range(1, len(cells), 2))
    if len(set(languages)) != len(languages):
        raise ValueError("第 %d 行表头的语言列重复" % line_no)
    return languages


def parse_markdown(path):
    """返回 [(产品号, 德语五点或 None, 法语五点或 None), ...]。"""
    with io.open(path, "r", encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    products = []
    cur_no = None
    cur_points = {"de": {}, "fr": {}}
    cur_languages = None
    seen_numbers = set()

    def flush():
        if cur_no is None:
            return
        if cur_languages is None:
            raise ValueError("产品 %s 没有可识别的五点表头" % cur_no)
        missing = [i for i in range(1, 6)
                   if any(i not in cur_points[lang] for lang in cur_languages)]
        if missing:
            raise ValueError("产品 %s 缺少第 %s 点" % (cur_no, missing))
        products.append((
            cur_no,
            [cur_points["de"][i] for i in range(1, 6)] if "de" in cur_languages else None,
            [cur_points["fr"][i] for i in range(1, 6)] if "fr" in cur_languages else None,
        ))

    for line_no, line in enumerate(lines, start=1):
        stripped = line.strip()
        head = PRODUCT_RE.match(stripped)
        if head:
            flush()
            cur_no = head.group(1)
            if cur_no in seen_numbers:
                raise ValueError("第 %d 行重复产品号 %s" % (line_no, cur_no))
            seen_numbers.add(cur_no)
            cur_points = {"de": {}, "fr": {}}
            cur_languages = None
            continue
        if cur_no is None or not stripped.startswith("|"):
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if cells[0] == "要点":
            languages = header_languages(cells, line_no)
            if cur_languages is not None and cur_languages != languages:
                raise ValueError("第 %d 行产品 %s 有不一致的语言表头" % (line_no, cur_no))
            cur_languages = languages
            continue
        if not ROW_RE.match(stripped):
            continue
        if cur_languages is None:
            raise ValueError("第 %d 行产品 %s 的要点前缺少语言表头" % (line_no, cur_no))
        if len(cells) != 1 + 2 * len(cur_languages):
            raise ValueError("第 %d 行产品 %s 的列数与表头不符" % (line_no, cur_no))
        i = int(cells[0])
        if not 1 <= i <= 5:
            raise ValueError("第 %d 行产品 %s 的要点编号 %d 不在 1~5 范围内" % (line_no, cur_no, i))
        for offset, lang in enumerate(cur_languages):
            point = cells[1 + 2 * offset]
            if not point:
                raise ValueError("产品 %s 第 %d 点存在空内容" % (cur_no, i))
            if i in cur_points[lang]:
                raise ValueError("第 %d 行产品 %s 的第 %d 点重复" % (line_no, cur_no, i))
            cur_points[lang][i] = point

    flush()
    if not products:
        raise ValueError("未从 %s 中解析到任何产品" % path)
    detect_mode(products)  # 提前拒绝同一文件内混用模式
    return products


def detect_mode(products):
    """由实际解析到的语言列决定输出模式。"""
    if not products:
        raise ValueError("没有可导出的产品")
    modes = set()
    for _no, de, fr in products:
        if de is not None and fr is not None:
            modes.add("de&fr")
        elif de is not None:
            modes.add("de")
        elif fr is not None:
            modes.add("fr")
        else:
            raise ValueError("产品缺少德语和法语五点")
    if len(modes) != 1:
        raise ValueError("同一文件混用了不同的语言模式")
    return modes.pop()


def write_xlsx(products, out_path):
    """单语每个产品一行；双语为德语整块 → 空两行 → 法语整块。"""
    mode = detect_mode(products)
    wb = Workbook()
    ws = wb.active
    ws.title = "fivepoints"

    n = len(products)
    fr_start = n + 3 if mode == "de&fr" else 1 if mode == "fr" else None
    for offset, (_no, de, fr) in enumerate(products):
        if de is not None:
            for col, point in enumerate(de, start=1):
                ws.cell(row=1 + offset, column=col).value = point
        if fr is not None:
            for col, point in enumerate(fr, start=1):
                ws.cell(row=fr_start + offset, column=col).value = point

    wb.save(out_path)
    return n, fr_start


def find_inputs():
    """只扫描脚本所在目录，不递归，也不读取外部路径。"""
    return sorted(
        (entry.path for entry in os.scandir(BASE)
         if entry.is_file() and not entry.is_symlink() and entry.name.lower().endswith(".md")),
        key=lambda path: os.path.basename(path).casefold(),
    )


def main(argv=None):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    parser = argparse.ArgumentParser(
        description="自动处理本目录的五点 Markdown，识别德语/法语/德法双语"
    )
    parser.parse_args(argv)

    inputs = find_inputs()
    if not inputs:
        print("[ERROR] %s 中没有 .md 文件" % BASE)
        return 1

    today = datetime.now().strftime("%Y%m%d")
    jobs = []
    seen_modes = {}
    for path in inputs:
        try:
            products = parse_markdown(path)
            mode = detect_mode(products)
        except (OSError, UnicodeError, ValueError) as exc:
            print("[ERROR] %s: %s" % (path, exc))
            return 1
        if mode in seen_modes:
            print("[ERROR] %s 与 %s 都是 %s 模式，当天输出文件会重名；请一次只放一份该模式的 Markdown"
                  % (seen_modes[mode], path, mode))
            return 1
        seen_modes[mode] = path
        jobs.append((path, products, mode, os.path.join(BASE, "%s_%s.xlsx" % (mode, today))))

    for path, products, mode, out_path in jobs:
        n, fr_start = write_xlsx(products, out_path)
        print("输入: %s" % path)
        print("输出: %s" % out_path)
        print("模式: %s | 产品数: %d" % (mode, n))
        if mode == "de&fr":
            print("布局: 德语第 1~%d 行；空两行；法语第 %d~%d 行" % (n, fr_start, fr_start + n - 1))
        else:
            print("布局: %s 第 1~%d 行（每行 5 列，无表头）" % (mode, n))
    return 0


if __name__ == "__main__":
    sys.exit(main())
