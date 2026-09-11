# -*- coding: utf-8 -*-
"""
antelope/common.py —— 各脚本共享的公共工具（配置加载 + JSON/分组/列工具）。

抽取动机：原来 7 个脚本各自重复实现了 zconfig 加载（importlib spec，因为
zconfig.constant.py 文件名含点无法直接 import）与 load_json / load_groups /
load_data_cols / load_col_scope 等小工具。这里统一收敛，行为不变。

用法：
    from common import zcfg, load_json, load_groups, ...
"""

import datetime
import getpass
import importlib.util
import json
import os
import re
import sys

from openpyxl.utils import column_index_from_string, get_column_letter

# ─────────────────────────────────────────────────────────────────────────── #
# 路径与 sys.path：antelope 与项目根都加入，便于 import zconfig / services
# ─────────────────────────────────────────────────────────────────────────── #
_ANTELOPE_DIR = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_ANTELOPE_DIR)
for _p in (_ANTELOPE_DIR, _ROOT):
    if _p not in sys.path:
        sys.path.insert(0, _p)


# ─────────────────────────────────────────────────────────────────────────── #
# zconfig 加载（zconfig.constant.py 文件名含点，无法用普通 import，故用 spec 加载）
# ─────────────────────────────────────────────────────────────────────────── #
def _load_zconfig():
    _p = os.path.join(_ANTELOPE_DIR, "zconfig.constant.py")
    _spec = importlib.util.spec_from_file_location("zconfig_constant", _p)
    _mod = importlib.util.module_from_spec(_spec)
    _spec.loader.exec_module(_mod)
    return _mod


zcfg = _load_zconfig()   # 模块级单例，所有脚本共享同一份配置


# ─────────────────────────────────────────────────────────────────────────── #
# 通用工具
# ─────────────────────────────────────────────────────────────────────────── #
def setup_utf8():
    """Windows 控制台以 UTF-8 输出（避免 GBK 报错）。"""
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass


# ─────────────────────────────────────────────────────────────────────────── #
# 统一日志：不再打印到控制台，全部写入 antelope/log/{当天日期}_atl_{操作用户}.log
# ─────────────────────────────────────────────────────────────────────────── #
def get_log_file() -> str:
    """返回日志文件路径：antelope/log/{YYYY-MM-DD}_atl_{用户}.log（目录不存在自动创建）。"""
    log_dir = os.path.join(_ANTELOPE_DIR, "log")
    os.makedirs(log_dir, exist_ok=True)
    date = datetime.date.today().strftime("%Y-%m-%d")
    try:
        user = getpass.getuser()
    except Exception:
        user = os.environ.get("USERNAME", "unknown")
    return os.path.join(log_dir, f"{date}_atl_{user}.log")


class _LogCleanWriter:
    """日志写出口：把不换行空格（\\xa0 / \\u2007 / \\u202f）统一替换为普通空格。

    模板表头常含 NBSP（如 "Piles nécessaires\\xa0?"），原样写入日志会显示异常，
    这里在写入端统一清洗，保证日志干净可读。
    """

    _MAP = str.maketrans({"\u00a0": " ", "\u2007": " ", "\u202f": " "})

    def __init__(self, fh):
        self._fh = fh

    def write(self, s):
        self._fh.write(str(s).translate(self._MAP))

    def flush(self):
        self._fh.flush()

    def __getattr__(self, name):
        return getattr(self._fh, name)


def setup_log() -> str:
    """把 stdout/stderr 统一重定向到当天日志文件（追加），控制台不再打印。

    所有 antelope 脚本在 main() 开头调用一次；同名同天多次运行/多个脚本
    （含 run_all 的子进程）都会追加进同一个日志文件，开头带运行分隔头。
    写入时自动清洗不换行空格（NBSP），日志中不会出现 \\xa0 等字符。
    """
    log_path = get_log_file()
    try:
        fh = open(log_path, "a", encoding="utf-8")
    except Exception as exc:
        # 打开日志失败时退回 UTF-8 控制台，避免流程不可见
        setup_utf8()
        print(f"⚠️ 无法写入日志文件 {log_path}: {exc}（退回控制台输出）")
        return log_path
    sys.stdout = _LogCleanWriter(fh)
    sys.stderr = _LogCleanWriter(fh)
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    argv = " ".join(sys.argv[1:])
    print(f"\n{'=' * 60}")
    print(f"[{now}] {os.path.basename(sys.argv[0])} {argv}".rstrip())
    print("=" * 60)
    return log_path


def load_json(path):
    """读取 JSON；utf-8-sig 兼容带 BOM 的人工编辑文件（如 PowerShell/记事本保存）。"""
    with open(path, "r", encoding="utf-8-sig") as f:
        return json.load(f)


def load_groups(path):
    """读取 groups JSON 的 groups 字段（分组行范围，实际行号）；缺失/失败返回 {}。"""
    if not path or not os.path.exists(path):
        return {}
    try:
        return load_json(path).get("groups") or {}
    except Exception:
        return {}


def load_data_cols(path):
    """读取 A 取数 data.json 中已覆盖的目标列（去重集合）。"""
    if not path or not os.path.exists(path):
        return set()
    try:
        data = load_json(path).get("data") or {}
    except Exception:
        return set()
    cols = set()
    for col_data in data.values():
        cols.update(str(c) for c in (col_data or {}).keys())
    return cols


def load_col_scope(diff_path):
    """读取 column_diff.json 的 only_in_completed 列号（待填列范围 col_scope）。"""
    if not path_exists(diff_path):
        return []
    try:
        diff = load_json(diff_path)
        only_in_completed = diff.get("by_col", {}).get("only_in_completed", [])
        return sorted(c["col"] for c in only_in_completed if "col" in c)
    except Exception:
        return []


def path_exists(path):
    return bool(path) and os.path.exists(path)


def parse_column_key(key):
    """配置键 → 列号 int：支持列字母（"AN"、"s"，大小写不敏感）与列号字符串（"40"）。

    识别不了（如 "_comment"、空串、乱写）返回 None，由调用方忽略该键。
    """
    s = str(key).strip().upper()
    if not s:
        return None
    if s.isdigit():
        return int(s)                      # 兼容旧写法：直接写列号
    if s.isalpha():
        try:
            return column_index_from_string(s)
        except Exception:
            return None
    return None


def column_letter(col):
    """列号 → 列字母（日志/报告里与配置文件保持同一写法，如 40 → "AN"）。"""
    try:
        return get_column_letter(int(col))
    except Exception:
        return str(col)


# ─────────────────────────────────────────────────────────────────────────── #
# 指定值的「序列种子」：S-0001 → S-0001, S-0002, …（整列连续递增，见 column_defaults.json）
# ─────────────────────────────────────────────────────────────────────────── #
def seq_seed(value):
    """判断指定值是否为可递增的序列种子 → (前缀, 起始数字, 数字位数)，否则 None。

    规则：值必须以数字结尾；前缀原样保留，数字位数用于决定左侧补零宽度。
        "S-0001" → ("S-", 1, 4)      "0001" → ("", 1, 4)      "A-1" → ("A-", 1, 1)
        "Coat-2024" → ("Coat-", 2024, 4)（同样是种子，不想要序列就写 sequence: false）
    """
    s = str(value)
    m = re.search(r"(\d+)$", s)
    if not m:
        return None
    digits = m.group(1)
    return s[: m.start(1)], int(digits), len(digits)


def make_sequence(seed, count):
    """按序列种子生成 count 个连续值；不是种子返回 None。

    数字递增、宽度不小于种子位数（左侧补零）；超出宽度时自然变长，不截断：
        ("S-0001", 3) → ["S-0001", "S-0002", "S-0003"]
        ("X-99", 3)   → ["X-99", "X-100", "X-101"]
    """
    info = seq_seed(seed)
    if info is None:
        return None
    prefix, start, width = info
    return [f"{prefix}{start + i:0{width}d}" for i in range(count)]


def normalize_sequence_flag(flag):
    """sequence 配置归一化 → False（非序列，默认）/ True（强制序列）/ None（auto 自动判断）。

    - 缺省、null、""、false/"false" → **False：整列同值（默认行为，不序列化）**；
    - true/"true" → True：强制序列（值必须以数字结尾，否则报错）；
    - "auto"/"default" → None：按值自动判断（值以数字结尾才序列化）；
    - 无法识别的写法 → False（安全默认：不序列化）。
    """
    if flag is None:
        return False
    if isinstance(flag, bool):
        return flag
    s = str(flag).strip().lower()
    if s in ("", "false", "no", "0", "off", "none", "fixed", "same", "n"):
        return False
    if s in ("true", "yes", "1", "on", "seq", "sequence", "y"):
        return True
    if s in ("auto", "default"):
        return None
    return False


def uncovered_cols(diff_path, data_path):
    """A 未覆盖的待填列：col_scope − A 已覆盖列。"""
    scope = load_col_scope(diff_path)
    a_cols = load_data_cols(data_path)
    return [c for c in scope if str(c) not in a_cols]


def load_ai_columns(completed_path, uncovered):
    """从「A 未覆盖列」中筛出有可选值(choices)的列，返回 [(col, header, choices)]。

    用于 ai_pick_attributes.py（AI 选值）；header 用于生成短标签。
    """
    completed = load_json(completed_path)
    by_col = {c["col"]: c for c in completed.get("columns", [])}
    result = []
    for col in uncovered:
        c = by_col.get(col)
        if not c:
            continue
        choices = c.get("choices") or []
        if choices:
            result.append((col, str(c.get("header", "")), [str(x) for x in choices]))
    return result


# ─────────────────────────────────────────────────────────────────────────── #
# 目标列（A 未覆盖 且 无可选值）取值：配置读取 / 值文件读取 / file>value 解析
#
# 概念：col_scope(C−B) 里 A 没映射、completed 也没有可选值(choices)的列，既没有取数
# 来源也没有 AI 选值依据 → 由 intermediate_tpl/column_defaults.json 人工指定
# 「一个值」或「一个多行值文件」，按整列连续循环填满（见 build_fill_framework.py）。
# ─────────────────────────────────────────────────────────────────────────── #
class ValueFileError(Exception):
    """值文件的数据错误（如空行）——属于必须人工修正的错误，不做兜底、直接报错退出。"""


def value_cycle_cols(diff_path, data_path, completed_path):
    """目标列 = 「A 未覆盖列」− 「有可选值(choices)的列」（本机制处理的对象）。

    A 未覆盖列见 uncovered_cols；有可选值的列由 ⑦ ai_pick_attributes.py 负责
    （且 ⑧ 已守门：不允许残留占位），故这里排除。
    """
    uncovered = uncovered_cols(diff_path, data_path)
    ai = {c for c, _, _ in load_ai_columns(completed_path, uncovered)}
    return [c for c in uncovered if c not in ai]


def load_column_defaults(path):
    """读取目标列取值配置，只取当前 ACTIVE_CATEGORY 标签下的部分。

    共享单文件结构（默认 intermediate_tpl/column_defaults.json）：
        {
          "_comment": ["…说明，键名以 _ 开头的一律忽略…"],
          "yass_fr_coat": {
            "S":  {"value": "Coat-001"},          ← 键 = **列字母**（Excel 列名，大小写不敏感）
            "AT": {"file": "keywords_coat.txt"},
            "AN": "Voir la description"           ← 字符串简写 = 只给 value
          }
        }

    键支持列字母（"S"/"AN"）与列号（"19"/"40"，旧写法兼容）；以 "_" 开头的键忽略。
    每项三个字段（推荐都写全，缺省按下面的默认值处理）：
        "value"    : 指定值；"" = 未配置（不写值）
        "file"     : 多行值文件；"" = 未配置（不写文件）
        "sequence" : 是否序列填充；**缺省 = false（非序列，整列同值）**，
                     true = 强制序列（值必须以数字结尾），"auto" = 值以数字结尾才序列化
    归一化为 {列号(str): {"value": str|None, "file": str|None, "sequence": bool|None}}：
    value/file 的空串与 null 一律归一为 None（表示未配置）；保留空项，
    便于区分「配置了但没值」与「未配置」两种警告。
    文件缺失/解析失败/找不到当前标签 → 返回 {}（全部列按未配置处理）。
    """
    if not path or not os.path.exists(path):
        return {}
    try:
        raw = load_json(path)
    except Exception:
        return {}
    if not isinstance(raw, dict):
        return {}
    section = raw.get(zcfg.ACTIVE_CATEGORY)
    if isinstance(section, dict):
        raw = section                      # 新版：取当前标签分段
    # 否则视为旧版扁平结构，raw 整体使用（值不是 dict/str 的项会被下面的归一化过滤）

    def _blank_to_none(v):
        """空串/null → None（=未配置）；其余转字符串。"""
        if v is None:
            return None
        s = str(v)
        return s if s.strip() else None

    result = {}
    for key, val in raw.items():
        col = parse_column_key(key)
        if col is None:
            continue                       # "_comment" 等非列键忽略
        if isinstance(val, str):
            result[str(col)] = {"value": _blank_to_none(val), "file": None, "sequence": False}
        elif isinstance(val, dict):
            result[str(col)] = {
                "value": _blank_to_none(val.get("value")),
                "file": _blank_to_none(val.get("file")),
                "sequence": normalize_sequence_flag(val.get("sequence")),
            }
    return result


def read_value_lines(path):
    """读取多行值文件 → [值...]（每行一个值，顺序即循环顺序）。

    - 编码 utf-8-sig 优先，失败退回 gbk（Windows 记事本另存 ANSI 的常见情况）；
    - 先把 CRLF / CR / 混用换行统一为 \\n 再切分 → 末尾换行不会多出一个值，
      也不会因残留 \\r 把一行拆成两行而误报空行；
    - **空行视为数据错误 → 抛 ValueFileError（带行号）**，不静默跳过（按需求：值文件不会有空值）；
    - 空文件（0 行）返回 []，由调用方决定是否退回 value；
    - 文件不存在 / 无法解码 → 原样抛出（FileNotFoundError / UnicodeDecodeError）。
    """
    data = None
    last_exc = None
    for enc in ("utf-8-sig", "gbk"):
        try:
            with open(path, "r", encoding=enc, newline="") as fh:
                data = fh.read()
            break
        except UnicodeDecodeError as exc:
            last_exc = exc
            continue
    if data is None:
        raise last_exc if last_exc else UnicodeDecodeError("unknown", b"", 0, 1, "decode failed")

    text = re.sub(r"\r+\n", "\n", data).replace("\r", "\n")   # CRLF/CRCRLF/单独CR 统一为 LF
    lines = text.split("\n")
    if lines and lines[-1] == "":
        lines.pop()                        # 文件末尾的换行不算一个值
    blanks = [i for i, ln in enumerate(lines, 1) if not ln.strip()]
    if blanks:
        raise ValueFileError(
            f"值文件存在空行（第 {blanks[:10]} 行{'...' if len(blanks) > 10 else ''}）: {path}"
            "—— 按约定值文件不允许空值，请删除空行后重跑。"
        )
    return lines


def resolve_column_values(entry, values_dir, log=print):
    """把一条配置解析为 (values, source, path)：file 优先，找不到退回 value。

    - file 可写绝对路径，或相对下列任一位置：`antelope/values/<类别>/`（默认目录）、
      `antelope/values/`（直接放在这一层也认）、仓库根、当前工作目录；
    - 文件不存在、读失败、为空（0 行）→ 退回 value（打印原因；没配 value 就保留占位）；
    - 空行等数据错误（ValueFileError）原样抛出，由调用方报错退出（不兜底）；
    - 都不行 → ([], None, None)。
    """
    value = (entry or {}).get("value")
    file_spec = (entry or {}).get("file")

    def _fallback_note():
        """退回说明：有指定值 → 改用指定值；没指定值 → 该列会保留占位。"""
        return "改用指定值" if value else "未配置指定值 → 该列保留占位"

    if file_spec:
        if os.path.isabs(file_spec):
            candidates = [file_spec]
        else:
            candidates = [
                os.path.join(values_dir or "", file_spec),           # antelope/values/<类别>/
                os.path.join(_ANTELOPE_DIR, "values", file_spec),    # antelope/values/
                os.path.join(_ROOT, file_spec),                      # 相对仓库根
                file_spec,                                           # 相对当前工作目录
            ]
        searched = []
        for c in candidates:                                          # 去重保序
            if c and c not in searched:
                searched.append(c)
        path = next((p for p in searched if path_exists(p)), None)
        if path is None:
            log(f"      ↳ ⚠️ 值文件不存在: {file_spec}")
            log(f"         已找: {searched}")
            hint = "" if value else (f"（把文件放到 {values_dir} 或 "
                                     f"{os.path.join(_ANTELOPE_DIR, 'values')} 下，或写相对仓库根的路径）")
            log(f"         → {_fallback_note()}{hint}")
        else:
            try:
                lines = read_value_lines(path)
            except ValueFileError:
                raise                          # 空行 = 数据错误，直接报错
            except Exception as exc:
                log(f"      ↳ ⚠️ 值文件读取失败（{type(exc).__name__}: {exc}）: {path}"
                    f" → {_fallback_note()}")
                lines = []
            if lines:
                return lines, "file", path
            log(f"      ↳ ⚠️ 值文件为空（0 行）: {path} → {_fallback_note()}")

    if value:
        return [value], "value", None
    return [], None, None


# ─────────────────────────────────────────────────────────────────────────── #
# 占位值（placeholder）工具
#
# 「有可选值(choices)的未覆盖列不允许残留占位」是同一条不变量，由两处共同守：
#   - ai_pick_attributes.py（⑦）：AI 选值写回 M 之后校验，残留 → exit 2；
#   - build_fill_framework.py（⑧）：生成 plan 之前守门，残留 → exit 2 且不写 plan。
# 判定实现集中在这里，避免两处口径不一致（曾出现 dataTemp / datatemp 混用）。
# ─────────────────────────────────────────────────────────────────────────── #
PLACEHOLDER_ALIASES = {"datatemp"}   # 历史遗留的小写拼写（大小写不敏感比较）


def is_placeholder(value, placeholder="dataTemp") -> bool:
    """判断单个值是否为占位值（大小写不敏感；兼容历史 dataTemp / datatemp 两种拼写）。"""
    s = str(value).strip().lower()
    return s == str(placeholder).strip().lower() or s in PLACEHOLDER_ALIASES


def is_placeholder_group(values, placeholder="dataTemp") -> bool:
    """判断某「组×列」是否为占位状态：缺失 / 空数组 / 全部为占位值。"""
    if values is None:
        return True
    if isinstance(values, str):
        return is_placeholder(values, placeholder)
    try:
        items = list(values)
    except TypeError:
        return is_placeholder(values, placeholder)
    if not items:
        return True
    return all(is_placeholder(v, placeholder) for v in items)


def find_residual_placeholders(m_data, groups, cols, placeholder="dataTemp"):
    """找出 cols 中仍残留占位的「组×列」（缺失 / 空数组 / 全占位 均算残留）。

    返回 {列号(int): {"groups": [组名...], "indexes": [组序号(1起)...], "count": n}}；
    无 cols 或全部干净 → {}。组 spec 非法的组按 update_m_data 的口径跳过。
    """
    result = {}
    want = sorted({int(c) for c in (cols or [])})
    if not want:
        return result
    for idx, (gname, spec) in enumerate((groups or {}).items(), 1):
        if str(spec).strip().count("&") != 1:
            continue
        gdata = (m_data or {}).get(str(gname)) or {}
        for col in want:
            values = gdata.get(str(col), gdata.get(col))
            if is_placeholder_group(values, placeholder):
                entry = result.setdefault(col, {"groups": [], "indexes": [], "count": 0})
                entry["groups"].append(str(gname))
                entry["indexes"].append(idx)
                entry["count"] += 1
    return result


# ─────────────────────────────────────────────────────────────────────────── #
# 单元格工具（复用 services/utils.py 的实现，但不触发 services/__init__.py 的
# 包级导入——那样会间接依赖 config.py；utils.py 本身只依赖 openpyxl）
# ─────────────────────────────────────────────────────────────────────────── #
_utils_spec = importlib.util.spec_from_file_location(
    "services_utils", os.path.join(_ROOT, "services", "utils.py")
)
_utils_mod = importlib.util.module_from_spec(_utils_spec)
_utils_spec.loader.exec_module(_utils_mod)
cell_has_fill = _utils_mod.cell_has_fill
