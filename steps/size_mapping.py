"""Apply the size mapping in place in the configured size column."""

from openpyxl.utils import get_column_letter

from core import PipelineContext, PipelineStep
from services import is_blank, cell_has_fill
from services.unmapped import record_unmapped


class SizeMappingStep(PipelineStep):
    name = "size_mapping"
    description = "Apply the size mapping in place"
    requires = ("insert_columns",)

    def run(self, ctx: PipelineContext) -> PipelineContext:
        cfg = self.config
        ws = ctx.worksheet
        mapping = cfg.load_size_mapping()
        if not mapping:
            size_col = get_column_letter(cfg.size_col)
            ctx.log(f"Size mapping is empty, skipping size column {size_col}")
            return ctx

        filled = 0
        cleared = 0
        skipped = 0
        for r in range(1, ws.max_row + 1):
            if cell_has_fill(ws.cell(row=r, column=cfg.col_a)):
                ws.cell(row=r, column=cfg.size_col).value = None
                cleared += 1
                continue
            val = ws.cell(row=r, column=cfg.size_col).value
            if is_blank(val):
                continue
            key = str(val).strip()
            if key in mapping:
                ws.cell(row=r, column=cfg.size_col).value = mapping[key]
                filled += 1
            else:
                skipped += 1
                record_unmapped("size", key, cfg.size_mapping_path, source="size_mapping")
        size_col = get_column_letter(cfg.size_col)
        ctx.log(f"Size column {size_col}: {filled} mapped, {cleared} cleared "
                f"(col A has fill), {skipped} no match (recorded)")
        return ctx
