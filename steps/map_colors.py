"""Map source color values into the mapped-color column."""

from openpyxl.utils import get_column_letter

from core import PipelineContext, PipelineStep
from services import is_blank, cell_has_fill
from services.unmapped import record_unmapped


class MapColorsStep(PipelineStep):
    name = "map_colors"
    description = "Map source colors into the mapped-color column"
    requires = ("insert_columns",)

    def run(self, ctx: PipelineContext) -> PipelineContext:
        cfg = self.config
        ws = ctx.worksheet
        mapping = cfg.load_color_mapping()
        if not mapping:
            target = get_column_letter(cfg.mapped_color_col)
            ctx.log(f"{cfg.color_mapping_path.name} is empty, skipping mapped-color column {target}")
            return ctx

        ci_mapping: dict[str, str] = {k.lower(): v for k, v in mapping.items()}
        filled = 0
        skipped = 0
        unmapped = 0
        for r in range(1, ws.max_row + 1):
            if cell_has_fill(ws.cell(row=r, column=cfg.col_a)):
                skipped += 1
                continue
            val = ws.cell(row=r, column=cfg.source_color_col).value
            if is_blank(val):
                continue
            key = str(val).strip().lower()
            if key in ci_mapping:
                ws.cell(row=r, column=cfg.mapped_color_col).value = ci_mapping[key]
                filled += 1
            else:
                unmapped += 1
                record_unmapped("color", str(val).strip(), cfg.color_mapping_path, source="map_colors")
        source = get_column_letter(cfg.source_color_col)
        target = get_column_letter(cfg.mapped_color_col)
        ctx.log(f"Colors {source} -> {target}: {filled} cells mapped, {skipped} rows skipped "
                f"(col A has fill), {unmapped} no match (recorded)")
        return ctx
