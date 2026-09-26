"""Calculate the configured price chain from the price-source text."""

from decimal import Decimal

from openpyxl.utils import get_column_letter

from core import PipelineContext, PipelineStep
from rules import extract_price_before_jpy
from services import is_blank, round_decimal


def _default_price_chain(base: Decimal, config) -> tuple[Decimal, Decimal, Decimal, Decimal]:
    """Return the price calculation chain with rounded intermediates."""
    base_price = base
    multiplied_price = Decimal(str(round_decimal(base * Decimal(config.price_multiplier))))
    subtracted_price = Decimal(str(round_decimal(multiplied_price - Decimal(config.price_subtract))))
    final_price = Decimal(str(round_decimal(subtracted_price + Decimal(config.price_add))))
    return base_price, multiplied_price, subtracted_price, final_price


class CalcPriceStep(PipelineStep):
    name = "calc_price"
    description = "Calculate the configured price chain"
    requires = ("mirror_category",)

    def run(self, ctx: PipelineContext) -> PipelineContext:
        cfg = self.config
        ws = ctx.worksheet
        calculator = cfg.price_calculator or (lambda value: _default_price_chain(value, cfg))
        filled = 0
        for r in range(1, ws.max_row + 1):
            text = ws.cell(row=r, column=cfg.price_source_col).value
            if is_blank(text):
                continue
            base = extract_price_before_jpy(str(text))
            if base is None:
                continue

            base_price, multiplied_price, subtracted_price, final_price = (
                round_decimal(value) for value in calculator(base)
            )

            ws.cell(row=r, column=cfg.base_price_col).value = base_price
            ws.cell(row=r, column=cfg.multiplied_price_col).value = multiplied_price
            ws.cell(row=r, column=cfg.subtracted_price_col).value = subtracted_price
            ws.cell(row=r, column=cfg.final_price_col).value = final_price
            filled += 1

        if filled:
            first = get_column_letter(cfg.base_price_col)
            last = get_column_letter(cfg.final_price_col)
            ctx.log(f"Price columns {first}-{last}: {filled} rows calculated")
        else:
            ctx.log("No JPY-anchored prices found, price columns skipped")
        return ctx
